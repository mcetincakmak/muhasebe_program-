"""Raporlar (satış/alış, KDV, brüt kâr, nakit) ve Excel'e aktarma."""
import io
import json
from collections import defaultdict
from datetime import date, datetime

from fastapi import APIRouter
from fastapi.responses import Response

from ..core import db
from ..servisler import stok
from .cari_hesap import CARI_KALEMLER, HESAP_KALEMLER, TURLER, cari_bakiyeleri, hesap_bakiyeleri
from ..core.ortak import hata

router = APIRouter()
AYLAR = ["Oca", "Şub", "Mar", "Nis", "May", "Haz", "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]


def _k(x):
    return round(float(x or 0) + 1e-9, 2)


def _aralik(bas, bit):
    bugun = date.today()
    try:
        b = date.fromisoformat(bas) if bas else date(bugun.year, 1, 1)
        e = date.fromisoformat(bit) if bit else bugun
    except ValueError:
        hata("Tarihleri YYYY-AA-GG biçiminde girin.")
    if b > e:
        hata("Başlangıç tarihi bitişten sonra olamaz.")
    return b.isoformat(), e.isoformat()


def rapor_verisi(bas, bit):
    bas, bit = _aralik(bas, bit)
    with db.islem() as con:
        satis = [dict(r) for r in con.execute(
            "SELECT id, fatura_no, tarih, fatura_turu, COALESCE(kur,1) AS kur, matrah*COALESCE(kur,1) AS matrah, "
            "kdv_toplam*COALESCE(kur,1) AS kdv_toplam, genel_toplam*COALESCE(kur,1) AS genel_toplam, "
            "COALESCE(tevkifat_toplam,0)*COALESCE(kur,1) AS tevkifat, cari_id, satirlar_json, "
            "json_extract(cari_json,'$.unvan') AS unvan FROM faturalar "
            "WHERE durum IN ('GONDERILDI','ONAYLANDI') AND tarih BETWEEN ? AND ?", (bas, bit))]
        alis = [dict(r) for r in con.execute(
            "SELECT id, fatura_no, tarih, tip, matrah*COALESCE(kur,1) AS matrah, kdv_toplam*COALESCE(kur,1) AS kdv_toplam, "
            "tutar*COALESCE(kur,1) AS tutar, COALESCE(tevkifat_toplam,0)*COALESCE(kur,1) AS tevkifat, cari_id, "
            "gonderen_unvan AS unvan "
            "FROM gelen_faturalar WHERE tarih BETWEEN ? AND ?", (bas, bit))]
        hareket = [dict(r) for r in con.execute(
            "SELECT tur, SUM(tutar) t FROM hareketler WHERE tarih BETWEEN ? AND ? GROUP BY tur", (bas, bit))]
        masraflar = [dict(r) for r in con.execute(
            "SELECT aciklama, SUM(tutar) t FROM hareketler WHERE tur='MASRAF' AND tarih BETWEEN ? AND ? "
            "GROUP BY aciklama ORDER BY t DESC LIMIT 10", (bas, bit))]
        cikislar = con.execute(
            "SELECT urun_id, -SUM(miktar) m FROM stok_hareketleri WHERE tur='SATIS' AND tarih BETWEEN ? AND ? "
            "GROUP BY urun_id", (bas, bit)).fetchall()
        maliyet = stok.ortalama_maliyetler(con, bit)

    # --- satış ve alış (iadeler düşülerek)
    satis_f = [f for f in satis if f["fatura_turu"] != "IADE"]
    alis_iade = [f for f in satis if f["fatura_turu"] == "IADE"]      # tedarikçiye kestiğimiz iade
    alis_f = [f for f in alis if f["tip"] != "IADE"]
    gelen_iade = [f for f in alis if f["tip"] == "IADE"]              # müşterinin bize kestiği iade

    def top(liste, alan):
        return _k(sum(f[alan] or 0 for f in liste))

    def matrah_alis(f):  # elle girilmiş faturada KDV bilinmiyorsa toplamı kullan
        return f["matrah"] if f["matrah"] is not None else f["tutar"]

    net_satis = top(satis_f, "matrah") - _k(sum(matrah_alis(f) for f in gelen_iade))
    net_alis = _k(sum(matrah_alis(f) for f in alis_f)) - top(alis_iade, "matrah")

    # --- aylık
    aylar = []
    y, a = int(bas[:4]), int(bas[5:7])
    while (y, a) <= (int(bit[:4]), int(bit[5:7])) and len(aylar) < 36:
        aylar.append(f"{y}-{a:02d}")
        y, a = (y + 1, 1) if a == 12 else (y, a + 1)
    aylik = {m: {"ay": m, "etiket": f"{AYLAR[int(m[5:]) - 1]} {m[2:4]}", "satis": 0.0, "alis": 0.0} for m in aylar}
    for f in satis_f:
        aylik[f["tarih"][:7]]["satis"] += f["matrah"] or 0
    for f in gelen_iade:
        aylik[f["tarih"][:7]]["satis"] -= matrah_alis(f) or 0
    for f in alis_f:
        aylik[f["tarih"][:7]]["alis"] += matrah_alis(f) or 0
    for f in alis_iade:
        aylik[f["tarih"][:7]]["alis"] -= f["matrah"] or 0

    # --- KDV
    hesaplanan = top(satis_f, "kdv_toplam")
    indirilecek = top([f for f in alis_f if f["kdv_toplam"] is not None], "kdv_toplam")
    iade_duzeltme_satis = top([f for f in gelen_iade if f["kdv_toplam"] is not None], "kdv_toplam")
    iade_duzeltme_alis = top(alis_iade, "kdv_toplam")
    kdv_bilinmeyen = sum(1 for f in alis if f["kdv_toplam"] is None)
    tevkif_satis = top(satis_f, "tevkifat")          # alıcının sorumlu sıfatıyla ödeyeceği kısım
    tevkif_alis = top(alis_f, "tevkifat")            # bizim sorumlu sıfatıyla (2 No.lu) ödeyeceğimiz
    kdv_net = _k((hesaplanan - iade_duzeltme_satis - tevkif_satis) - (indirilecek - iade_duzeltme_alis))

    # --- brüt kâr (stoklu ürünlerin maliyeti)
    smm, maliyetsiz = 0.0, 0
    for c in cikislar:
        if c["urun_id"] in maliyet:
            smm += c["m"] * maliyet[c["urun_id"]]
        else:
            maliyetsiz += 1
    hk = {h["tur"]: _k(h["t"]) for h in hareket}
    brut_kar = _k(net_satis - smm)

    # --- en çoklar
    musteri, tedarikci, urun = defaultdict(float), defaultdict(float), defaultdict(lambda: [0.0, 0.0])
    for f in satis_f:
        musteri[f["unvan"] or "?"] += f["matrah"] or 0
        for s in json.loads(f["satirlar_json"] or "[]"):
            urun[s.get("ad") or "?"][0] += float(s.get("miktar") or 0)
            urun[s.get("ad") or "?"][1] += float(s.get("matrah") or 0) * f["kur"]
    for f in alis_f:
        tedarikci[f["unvan"] or "?"] += matrah_alis(f) or 0

    def ilk5(d):
        return [{"ad": k, "tutar": _k(v)} for k, v in sorted(d.items(), key=lambda x: -x[1])[:5]]

    return {
        "bas": bas, "bit": bit,
        "ozet": {"net_satis": _k(net_satis), "net_alis": _k(net_alis), "smm": _k(smm), "brut_kar": brut_kar,
                 "masraf": hk.get("MASRAF", 0), "diger_gelir": hk.get("GELIR", 0),
                 "faaliyet_sonucu": _k(brut_kar - hk.get("MASRAF", 0) + hk.get("GELIR", 0)),
                 "satis_adet": len(satis_f), "alis_adet": len(alis_f), "maliyetsiz_urun": maliyetsiz},
        "aylik": [{**v, "satis": _k(v["satis"]), "alis": _k(v["alis"])} for v in aylik.values()],
        "kdv": {"hesaplanan": hesaplanan, "musteri_iadesi": iade_duzeltme_satis, "indirilecek": indirilecek,
                "tedarikci_iadesi": iade_duzeltme_alis, "net": kdv_net, "kdv_bilinmeyen_alis": kdv_bilinmeyen,
                "tevkif_satis": tevkif_satis, "tevkif_alis": tevkif_alis},
        "nakit": {"tahsilat": hk.get("TAHSILAT", 0), "odeme": hk.get("ODEME", 0), "masraf": hk.get("MASRAF", 0),
                  "gelir": hk.get("GELIR", 0)},
        "masraflar": [{"ad": m["aciklama"] or "?", "tutar": _k(m["t"])} for m in masraflar],
        "musteriler": ilk5(musteri), "tedarikciler": ilk5(tedarikci),
        "urunler": [{"ad": k, "miktar": round(v[0], 3), "tutar": _k(v[1])}
                    for k, v in sorted(urun.items(), key=lambda x: -x[1][1])[:5]],
    }


@router.get("/api/rapor")
def rapor(bas: str = "", bit: str = ""):
    return rapor_verisi(bas, bit)


# ------------------------------------------------------------------ Excel
def excel(bas, bit):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    r = rapor_verisi(bas, bit)
    bas, bit = r["bas"], r["bit"]
    wb = Workbook()
    F = Font(name="Arial", size=10)
    FB = Font(name="Arial", size=10, bold=True)
    FH = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    DOLGU = PatternFill("solid", fgColor="1E5B4A")
    CIZGI = Border(top=Side(style="thin", color="1D2521"))
    PARA = '#,##0.00;[Red]-#,##0.00;"-"'

    def sayfa(ws, basliklar, satirlar, para_sutunlari=(), toplam_sutunlari=(), genislik=None):
        ws.append(basliklar)
        for c in ws[1]:
            c.font, c.fill, c.alignment = FH, DOLGU, Alignment(vertical="center")
        for s in satirlar:
            ws.append(s)
        son = ws.max_row
        for row in ws.iter_rows(min_row=2, max_row=son):
            for c in row:
                c.font = F
                # Formül enjeksiyonu: tedarikçi adı gibi dış veriler "=" ile başlıyorsa formül değil metin olsun
                if c.data_type == "f":
                    c.data_type = "s"
                if c.column in para_sutunlari:
                    c.number_format = PARA
        if toplam_sutunlari and son >= 2:
            t = son + 2
            ws.cell(t, 1, "Toplam").font = FB
            for col in toplam_sutunlari:
                h = ws.cell(t, col, f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{son})")
                h.font, h.number_format, h.border = FB, PARA, CIZGI
        for i, g in enumerate(genislik or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = g
        ws.freeze_panes = "A2"

    # Özet
    ws = wb.active
    ws.title = "Özet"
    a = db.ayarlar_hepsi()
    o, k = r["ozet"], r["kdv"]
    satirlar = [
        [a.get("firma_unvan") or "Firma"], [f"Dönem: {bas} – {bit}"], [f"Oluşturulma: {datetime.now():%d.%m.%Y %H:%M}"], [],
        ["Satış ve alış (KDV hariç)", "Tutar (TL)"],
        ["Net satış (iadeler düşülmüş)", o["net_satis"]], ["Net alış (iadeler düşülmüş)", o["net_alis"]],
        ["Satılan stoklu malın maliyeti", o["smm"]], ["Brüt kâr", "=B6-B8"],
        ["Masraflar", o["masraf"]], ["Diğer gelirler", o["diger_gelir"]], ["Faaliyet sonucu (yaklaşık)", "=B9-B10+B11"], [],
        ["KDV", "Tutar (TL)"],
        ["Hesaplanan KDV (satışlar)", k["hesaplanan"]], ["Müşteri iadelerinden düzeltme", k["musteri_iadesi"]],
        ["İndirilecek KDV (alışlar)", k["indirilecek"]], ["Tedarikçiye iadelerden düzeltme", k["tedarikci_iadesi"]],
        ["Tevkif edilen KDV (satışlarda alıcı öder)", k["tevkif_satis"]],
        ["Fark (artı: ödenecek, eksi: devreden)", "=(B15-B16-B19)-(B17-B18)"],
        ["Sorumlu sıfatıyla ödenecek KDV (alışlardaki tevkifat, 2 No.lu)", k["tevkif_alis"]], [],
        ["Uyarılar"],
        [f"KDV'si bilinmeyen (elle toplam girilmiş) alış faturası: {k['kdv_bilinmeyen_alis']}"],
        [f"Maliyeti bilinmeyen satılan ürün: {o['maliyetsiz_urun']} (brüt kâra maliyetsiz yansıdı)"],
        ["Bu rapor bilgi amaçlıdır; KDV beyannamesi için mali müşavirinizin kayıtlarını esas alın."],
    ]
    for s in satirlar:
        ws.append(s)
    for row in ws.iter_rows():
        for c in row:
            c.font = F
            if c.column == 2 and c.row > 5:
                c.number_format = PARA
    for rr in (1, 5, 14, 23):
        ws.cell(rr, 1).font = FB
        ws.cell(rr, 2).font = FB
    for rr in (9, 12, 20):
        ws.cell(rr, 1).font = FB
        ws.cell(rr, 2).font = FB
        ws.cell(rr, 2).border = CIZGI
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18

    with db.islem() as con:
        # Satış faturaları
        rows = [[f["tarih"], f["fatura_no"], "İade" if f["fatura_turu"] == "IADE" else "Satış",
                 {"EFATURA": "e-Fatura", "EARSIV": "e-Arşiv"}.get(f["tip"], ""), f["unvan"], f["vkn"],
                 f["para_birimi"] or "TRY", f["kur"] or 1,
                 _k((f["matrah"] or 0) * (f["kur"] or 1)), _k((f["kdv_toplam"] or 0) * (f["kur"] or 1)),
                 _k((f["tevkifat_toplam"] or 0) * (f["kur"] or 1)), _k((f["genel_toplam"] or 0) * (f["kur"] or 1)), f["durum"]]
                for f in con.execute("SELECT f.*, json_extract(cari_json,'$.unvan') unvan, json_extract(cari_json,'$.vkn') vkn "
                                     "FROM faturalar f WHERE durum IN ('GONDERILDI','ONAYLANDI') AND tarih BETWEEN ? AND ? "
                                     "ORDER BY tarih, fatura_no", (bas, bit))]
        sayfa(wb.create_sheet("Satış faturaları"),
              ["Tarih", "Fatura no", "Tür", "Belge", "Cari", "VKN/TCKN", "Para", "Kur", "KDV hariç (TL)", "KDV (TL)",
               "Tevkifat (TL)", "Ödenecek (TL)", "Durum"],
              rows, (9, 10, 11, 12), (9, 10, 11, 12), [11, 18, 8, 10, 36, 13, 6, 9, 15, 12, 13, 15, 12])
        # Alış faturaları
        rows = [[g["tarih"], g["fatura_no"], "İade" if g["tip"] == "IADE" else "Alış",
                 "Elle" if g["kaynak"] == "manuel" else "e-Fatura", g["gonderen_unvan"], g["gonderen_vkn"],
                 g["para_birimi"] or "TRY", g["kur"] or 1,
                 _k(g["matrah"] * (g["kur"] or 1)) if g["matrah"] is not None else None,
                 _k(g["kdv_toplam"] * (g["kur"] or 1)) if g["kdv_toplam"] is not None else None,
                 _k((g["tevkifat_toplam"] or 0) * (g["kur"] or 1)), _k((g["tutar"] or 0) * (g["kur"] or 1)),
                 ", ".join(json.loads(g["irsaliye_nolar"] or "[]"))]
                for g in con.execute("SELECT * FROM gelen_faturalar WHERE tarih BETWEEN ? AND ? ORDER BY tarih", (bas, bit))]
        sayfa(wb.create_sheet("Alış faturaları"),
              ["Tarih", "Fatura no", "Tür", "Kaynak", "Tedarikçi", "VKN/TCKN", "Para", "Kur", "KDV hariç (TL)", "KDV (TL)",
               "Tevkifat (TL)", "Ödenecek (TL)", "İrsaliye no"],
              rows, (9, 10, 11, 12), (9, 10, 11, 12), [11, 18, 8, 10, 36, 13, 6, 9, 15, 12, 13, 15, 22])
        # Cariler
        b = cari_bakiyeleri(con)
        rows = [[c["unvan"], c["vkn"], {"musteri": "Müşteri", "tedarikci": "Tedarikçi", "ikisi": "İkisi"}.get(c["tur"], ""),
                 c["telefon"], c["eposta"], max(b.get(c["id"], 0), 0), max(-b.get(c["id"], 0), 0)]
                for c in con.execute("SELECT * FROM cariler ORDER BY unvan COLLATE NOCASE")]
        sayfa(wb.create_sheet("Cari bakiyeler"),
              ["Ünvan", "VKN/TCKN", "Tür", "Telefon", "E-posta", "Size borçlu", "Siz borçlusunuz"],
              rows, (6, 7), (6, 7), [38, 13, 11, 15, 26, 15, 15])
        # Cari hareketler
        rows = [[h["tarih"], h["unvan"], h["islem"], h["belge_no"], h["aciklama"], h["borc"], h["alacak"]]
                for h in con.execute(f"SELECT k.*, c.unvan FROM ({CARI_KALEMLER}) k JOIN cariler c ON c.id=k.cari_id "
                                     "WHERE k.tarih BETWEEN ? AND ? ORDER BY k.tarih, c.unvan", (bas, bit))]
        sayfa(wb.create_sheet("Cari hareketler"), ["Tarih", "Cari", "İşlem", "Belge", "Açıklama", "Borç", "Alacak"],
              rows, (6, 7), (6, 7), [11, 36, 26, 18, 30, 14, 14])
        # Stok
        m, mal = stok.miktarlar(con), stok.ortalama_maliyetler(con)
        rows = []
        for u in con.execute("SELECT * FROM urunler WHERE aktif=1 AND stok_takibi=1 ORDER BY ad COLLATE NOCASE"):
            miktar = m.get(u["id"], 0)
            rows.append([u["kod"], u["ad"], miktar, u["kritik"] or 0, mal.get(u["id"]), None,
                         "Kritik" if (u["kritik"] or 0) > 0 and miktar <= u["kritik"] else ""])
        ws2 = wb.create_sheet("Stok")
        sayfa(ws2, ["Kod", "Ürün", "Miktar", "Kritik seviye", "Ort. maliyet", "Stok değeri", "Durum"],
              rows, (5, 6), (), [12, 38, 10, 12, 14, 15, 10])
        for i in range(2, ws2.max_row + 1):
            ws2.cell(i, 6, f'=IF(AND(ISNUMBER(C{i}),ISNUMBER(E{i}),C{i}>0),C{i}*E{i},0)').number_format = PARA
            ws2.cell(i, 6).font = F
        if rows:
            t = ws2.max_row + 2
            ws2.cell(t, 1, "Toplam stok değeri").font = FB
            h = ws2.cell(t, 6, f"=SUM(F2:F{t - 2})")
            h.font, h.number_format, h.border = FB, PARA, CIZGI
        # Kasa / banka
        hb = hesap_bakiyeleri(con)
        ws3 = wb.create_sheet("Kasa ve banka")
        rows = []
        for hs in con.execute("SELECT * FROM hesaplar WHERE aktif=1"):
            for x in con.execute(f"SELECT * FROM ({HESAP_KALEMLER}) WHERE tarih BETWEEN :bas AND :bit ORDER BY tarih, id",
                                 {"hid": hs["id"], "bas": bas, "bit": bit}):
                rows.append([x["tarih"], hs["ad"], TURLER.get(x["tur"], x["tur"]), x["cari_unvan"], x["odeme_sekli"],
                             x["belge_no"], x["aciklama"], x["giris"], x["cikis"]])
        rows.sort(key=lambda z: z[0])
        sayfa(ws3, ["Tarih", "Hesap", "İşlem", "Cari", "Şekil", "Belge", "Açıklama", "Giriş", "Çıkış"],
              rows, (8, 9), (8, 9), [11, 18, 22, 30, 12, 14, 26, 14, 14])
        t = ws3.max_row + 2
        ws3.cell(t, 1, "Güncel bakiyeler").font = FB
        for hs in con.execute("SELECT * FROM hesaplar WHERE aktif=1"):
            t += 1
            ws3.cell(t, 2, hs["ad"]).font = F
            c = ws3.cell(t, 9, hb.get(hs["id"], 0))
            c.font, c.number_format = F, PARA

    # Aylık
    ws4 = wb.create_sheet("Aylık")
    sayfa(ws4, ["Ay", "Net satış", "Net alış", "Fark"], [[x["etiket"], x["satis"], x["alis"], None] for x in r["aylik"]],
          (2, 3, 4), (2, 3), [10, 16, 16, 16])
    for i in range(2, len(r["aylik"]) + 2):
        ws4.cell(i, 4, f"=B{i}-C{i}").number_format = PARA
        ws4.cell(i, 4).font = F

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), bas, bit


@router.get("/api/disa-aktar")
def disa_aktar(bas: str = "", bit: str = ""):
    veri, b, e = excel(bas, bit)
    return Response(veri, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="rapor_{b}_{e}.xlsx"'})
