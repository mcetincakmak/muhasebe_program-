"""Uçtan uca testler: python -m pytest tests  (deneme modunda, QNB'ye bağlanmaz)"""
import base64
import os
import tempfile

import pytest

_gecici = tempfile.mkdtemp()
os.environ["FATURA_VERI"] = _gecici
os.environ["FATURA_ZAMANLAYICI"] = "0"

from fastapi.testclient import TestClient  # noqa: E402

from app.core import db  # noqa: E402
from app.main import app  # noqa: E402
from app.servisler.ubl import fatura_xml, hesapla, xml_ozet  # noqa: E402


@pytest.fixture(scope="module")
def c():
    istemci = TestClient(app)
    r = istemci.post("/api/kurulum", json={"kullanici_adi": "yonetici", "sifre": "gizli1234", "firma_unvan": "Test Ltd. Şti."})
    assert r.status_code == 200, r.text
    istemci.put("/api/ayarlar", json={"firma_unvan": "Test Ltd. Şti.", "firma_vkn": "4810012345",
                                      "firma_vergi_dairesi": "Burdur", "firma_il": "Burdur"})
    return istemci


def test_hesap_iskonto_kdv_tevkifat():
    h = hesapla([{"ad": "Temizlik", "miktar": 1, "birim_fiyat": 1000, "kdv": 20, "iskonto": 10, "tevkifat_kod": "612"},
                 {"ad": "Malzeme", "miktar": 3, "birim_fiyat": 33.335, "kdv": 10}])
    assert h["matrah"] == 900 + 100.01
    assert h["kdv_toplam"] == 180 + 10.0
    assert h["tevkifat_toplam"] == 162.0
    assert h["genel_toplam"] == round(1000.01 + 190 - 162, 2)


def test_xml_dovizli_tevkifatli_geri_okunur():
    x = fatura_xml({"fatura_no": "ABC2026000000001", "uuid": "u-1", "tip": "EFATURA", "tarih": "2026-09-01",
                    "para_birimi": "USD", "kur": 40, "satirlar": [{"ad": "a", "miktar": 1, "birim_fiyat": 100, "kdv": 20,
                                                                  "tevkifat_kod": "604"}]},
                   {"unvan": "S", "vkn": "1234567890"}, {"unvan": "A", "vkn": "1111111111"})
    o = xml_ozet(x)
    assert o["tip"] == "TEVKIFAT" and o["para_birimi"] == "USD" and o["kur"] == 40
    assert o["tutar"] == 110.0 and o["tevkifat_toplam"] == 10.0


def test_xml_bombasi_reddedilir():
    bomba = '<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]><Invoice>&b;</Invoice>'
    with pytest.raises(Exception):
        xml_ozet(bomba)


def test_oturumsuz_erisim_yok():
    assert TestClient(app).get("/api/faturalar").status_code == 401


def test_farkli_kaynaktan_istek_reddedilir(c):
    r = c.post("/api/cariler", json={"unvan": "X"}, headers={"Origin": "https://kotu-site.example"})
    assert r.status_code == 403


def test_gecersiz_kdv_ve_miktar_reddedilir(c):
    cari = c.post("/api/cariler", json={"unvan": "Müşteri", "vkn": "1234567890"}).json()
    r = c.post("/api/faturalar", json={"cari_id": cari["id"], "satirlar": [{"ad": "x", "miktar": 1, "birim_fiyat": 1, "kdv": 7}]})
    assert r.status_code == 400
    r = c.post("/api/faturalar", json={"cari_id": cari["id"], "satirlar": [{"ad": "x", "miktar": "abc", "birim_fiyat": 1}]})
    assert r.status_code == 400


def test_satis_bakiye_stok_ve_rapor(c):
    urun = c.post("/api/urunler", json={"ad": "Kağıt", "kod": "KGT", "fiyat": 100, "kritik": 5}).json()
    c.post(f"/api/urunler/{urun['id']}/stok-duzelt", json={"tur": "ACILIS", "miktar": 10, "birim_fiyat": 60})
    cari = c.post("/api/cariler", json={"unvan": "Alıcı A.Ş.", "vkn": "2222222222"}).json()
    f = c.post("/api/faturalar", json={"cari_id": cari["id"], "satirlar": [
        {"ad": "Kağıt", "miktar": 6, "birim": "C62", "birim_fiyat": 100, "kdv": 20}]}).json()
    g = c.post(f"/api/faturalar/{f['id']}/gonder").json()
    assert g["durum"] in ("GONDERILDI", "ONAYLANDI") and g["fatura_no"]
    assert c.get(f"/api/cariler/{cari['id']}").json()["bakiye"] == 720.0
    stok = {u["id"]: u for u in c.get("/api/stok").json()}[urun["id"]]
    assert stok["miktar"] == 4 and stok["kritik_durum"]
    kasa = c.get("/api/hesaplar").json()[0]
    c.post("/api/hareketler", json={"tur": "TAHSILAT", "cari_id": cari["id"], "hesap_id": kasa["id"],
                                    "tarih": f["tarih"], "tutar": 720})
    assert c.get(f"/api/cariler/{cari['id']}").json()["bakiye"] == 0
    r = c.get("/api/rapor").json()
    assert r["ozet"]["smm"] == 360.0


def test_irsaliye_fotograf_ve_eslestirme(c):
    yol = os.path.join(os.path.dirname(__file__), "irsaliye_ornek.jpg")
    if not os.path.exists(yol):
        pytest.skip("örnek irsaliye görüntüsü yok")
    foto = "data:image/jpeg;base64," + base64.b64encode(open(yol, "rb").read()).decode()
    i = c.post("/api/irsaliyeler/yukle", json={"fotolar": [foto]}).json()
    assert i["irsaliye_no"] == "GKI2026000000031"
    c.post("/api/tarama")
    assert c.get(f"/api/irsaliyeler/{i['id']}").json()["durum"] == "ESLESTI"
    # aynı irsaliye ikinci kez kaydedilemez
    i2 = c.post("/api/irsaliyeler/yukle", json={"fotolar": [foto]}).json()
    assert i2["durum"] == "KONTROL" and "zaten kayıtlı" in i2["okuma"]["_hata"]


def test_excel_formul_enjeksiyonu(c):
    c.post("/api/cariler", json={"unvan": '=HYPERLINK("http://x","tikla")', "vkn": "3333333333"})
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(c.get("/api/disa-aktar").content))
    ws = wb["Cari bakiyeler"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value and str(row[0].value).startswith("=HYPERLINK"):
            assert row[0].data_type == "s"
            break
    else:
        pytest.fail("test carisi bulunamadı")


def test_giris_denemesi_kilidi():
    t = TestClient(app)
    for _ in range(5):
        t.post("/api/giris", json={"kullanici_adi": "yonetici", "sifre": "yanlis"})
    assert t.post("/api/giris", json={"kullanici_adi": "yonetici", "sifre": "gizli1234"}).status_code == 429


def _giris(kadi, sifre):
    t = TestClient(app)
    r = t.post("/api/giris", json={"kullanici_adi": kadi, "sifre": sifre})
    assert r.status_code == 200, r.text
    return t


def test_firmalar_birbirinden_yalitilmis(c):
    ikinci = c.post("/api/firmalar", json={"unvan": "İkinci Firma"}).json()
    ilk_firma = c.get("/api/ben").json()["firma"]["id"]
    c.post("/api/cariler", json={"unvan": "Sadece Birinci Firmada", "vkn": "4444444444"})
    c.post("/api/firma-sec", json={"firma_id": ikinci["id"]})
    assert c.get("/api/ben").json()["firma"]["unvan"] == "İkinci Firma"
    assert not any(x["unvan"] == "Sadece Birinci Firmada" for x in c.get("/api/cariler").json())
    c.post("/api/firma-sec", json={"firma_id": ilk_firma})
    assert any(x["unvan"] == "Sadece Birinci Firmada" for x in c.get("/api/cariler").json())


def test_roller_ve_yetkiler(c):
    firma_id = c.get("/api/ben").json()["firma"]["id"]
    r = c.post("/api/kullanicilar", json={"kullanici_adi": "depo", "sifre": "depo12345",
                                           "yetkiler": [{"firma_id": firma_id, "rol": "personel"}]})
    assert r.status_code == 200, r.text
    p = _giris("depo", "depo12345")
    assert p.get("/api/irsaliyeler").status_code == 200          # irsaliye görebilir
    assert p.get("/api/faturalar").status_code == 403            # faturaları göremez
    assert p.get("/api/kullanicilar").status_code == 403         # kullanıcı yönetemez
    assert p.post("/api/cariler", json={"unvan": "X"}).status_code == 403
    c.post("/api/kullanicilar", json={"kullanici_adi": "musavir", "sifre": "musavir123",
                                      "yetkiler": [{"firma_id": firma_id, "rol": "musavir"}]})
    m = _giris("musavir", "musavir123")
    assert m.get("/api/faturalar").status_code == 200 and m.get("/api/rapor").status_code == 200
    assert m.post("/api/cariler", json={"unvan": "X"}).status_code == 403   # salt okuma


def test_baska_firmaya_gecilemez(c):
    p = _giris("depo", "depo12345")
    assert p.post("/api/firma-sec", json={"firma_id": 999}).status_code == 403


def test_son_yonetici_pasiflestirilemez(c):
    ben = c.get("/api/ben").json()["kullanici"]
    r = c.put(f"/api/kullanicilar/{ben['id']}", json={"ad": "Y", "aktif": False, "sistem_yoneticisi": True})
    assert r.status_code == 400


def test_lisans(c):
    from app.core import lisans
    d = c.get("/api/lisans").json()
    assert d["tur"] == "deneme" and d["gecerli"]
    assert c.post("/api/lisans", json={"anahtar": "sahte.anahtar"}).status_code == 400
    ozel_yol = os.environ.get("LISANS_OZEL_ANAHTAR")
    if not ozel_yol or not os.path.exists(ozel_yol):
        pytest.skip("özel anahtar yok")
    import subprocess
    import sys
    anahtar = subprocess.run([sys.executable, "tools/lisans_uret.py", "--anahtar", ozel_yol, "--musteri", "Test",
                              "--bitis", "2099-01-01", "--firma", "5", "--kullanici", "10",
                              "--makine", lisans.makine_kodu()], capture_output=True, text=True).stdout.strip()
    d = c.post("/api/lisans", json={"anahtar": anahtar}).json()
    assert d["tur"] == "lisans" and d["firma"] == 5 and d["gecerli"]
    baska = subprocess.run([sys.executable, "tools/lisans_uret.py", "--anahtar", ozel_yol, "--musteri", "T",
                            "--bitis", "2099-01-01", "--makine", "AAAA-BBBB-CCCC-DDDD"],
                           capture_output=True, text=True).stdout.strip()
    assert "başka bir bilgisayar" in c.post("/api/lisans", json={"anahtar": baska}).json()["detail"]


def test_lisans_bitince_salt_okuma(c, monkeypatch):
    from app.core import lisans
    monkeypatch.setattr(lisans, "durum", lambda: {"tur": "deneme", "gecerli": False, "firma": 1, "kullanici": 1})
    assert c.get("/api/cariler").status_code == 200
    assert c.post("/api/cariler", json={"unvan": "Yeni"}).status_code == 403
    assert c.get("/api/disa-aktar").status_code == 200


def test_bozuk_gelen_belge_taramayi_durdurmaz(c, monkeypatch):
    from app.entegrator import deneme
    # Yeni (kaydedilecek) bir belgeden SONRA gelen bozuk belge: kayıt işlemi açıkken günlüğe yazmaya çalışmamalı
    yeni = fatura_xml({"fatura_no": "YNI2026000000001", "uuid": "bozuk-testi-yeni", "tip": "EFATURA", "tarih": "2026-09-01",
                       "satirlar": [{"ad": "Toner", "miktar": 1, "birim_fiyat": 500, "kdv": 20}]},
                      {"unvan": "Yeni Tedarikçi Ltd.", "vkn": "5555555555"}, {"unvan": "Test", "vkn": "4810012345"})
    monkeypatch.setattr(deneme.Deneme, "gelen_faturalar",
                        lambda self, son=0: ([("YENI-1", yeni), ("BOZUK-1", "<bozuk")], son))
    monkeypatch.setattr(db, "_baglan", _kisa_zaman_asimli(db._baglan))
    r = c.post("/api/tarama")
    assert r.status_code == 200, r.text
    assert r.json()["eklenen"] == 1
    assert any("BOZUK-1" in g["mesaj"] for g in c.get("/api/gunluk").json())


def _kisa_zaman_asimli(baglan):
    """Kilitlenmeyi 30 sn beklemek yerine hemen hata olarak görmek için."""
    def sar(yol):
        con = baglan(yol)
        con.execute("PRAGMA busy_timeout = 200")
        return con
    return sar


def test_lisans_siniri_yeniden_etkinlestirmede_de_gecerli(c):
    kullanicilar = {k["kullanici_adi"]: k for k in c.get("/api/kullanicilar").json()}
    musavir = kullanicilar["musavir"]
    assert c.put(f"/api/kullanicilar/{musavir['id']}", json={"ad": "M", "aktif": False}).status_code == 200
    assert c.post("/api/kullanicilar", json={"kullanici_adi": "yedek", "sifre": "yedek1234"}).status_code == 200
    r = c.put(f"/api/kullanicilar/{musavir['id']}", json={"ad": "M", "aktif": True})
    assert r.status_code == 400 and "kullanıcı" in r.json()["detail"]

    aktif = c.get("/api/ben").json()["firma"]["id"]
    diger = next(f for f in c.get("/api/firmalar").json() if f["id"] != aktif)
    assert c.put(f"/api/firmalar/{diger['id']}", json={"aktif": False}).status_code == 200
    assert c.post("/api/firmalar", json={"unvan": "Üçüncü Firma"}).status_code == 200
    r = c.put(f"/api/firmalar/{diger['id']}", json={"aktif": True})
    assert r.status_code == 400 and "firma" in r.json()["detail"]
